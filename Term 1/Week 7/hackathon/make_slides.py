"""Build slides.pptx for the 5-10 minute presentation. Run: .venv/bin/python make_slides.py"""
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

GREEN = RGBColor(0x1F, 0x5C, 0x3A)
DARK = RGBColor(0x22, 0x22, 0x22)
GREY = RGBColor(0x66, 0x66, 0x66)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def text(slide, x, y, w, h, lines, size=22, color=DARK, bold_first=False):
    tf = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.font.bold = bold_first and i == 0
        p.space_after = Pt(10)
    return tf


def slide(title, bullets, image=None, note=None):
    s = prs.slides.add_slide(BLANK)
    text(s, 0.6, 0.4, 12, 1, [title], size=36, color=GREEN, bold_first=True)
    width = 6.6 if image else 12
    text(s, 0.6, 1.5, width, 5.2, ["• " + b for b in bullets], size=22)
    if image:
        s.shapes.add_picture(image, Inches(7.6), Inches(1.5), height=Inches(5.3))
    if note:
        text(s, 0.6, 6.85, 12, 0.5, [note], size=14, color=GREY)
    return s


s = prs.slides.add_slide(BLANK)
text(s, 0.8, 2.2, 11.5, 1.5, ["Knotweed Checker"], size=54, color=GREEN, bold_first=True)
text(s, 0.8, 3.5, 11.5, 1.5, ["Spot invasive Japanese knotweed from a phone photo",
                              "AI for Good, Hackathon 7: Eyes on Earth · Roboflow · SDG 15"], size=26)

slide("The problem", [
    "Japanese knotweed pushes out native plants, can damage river banks and paving, and regrows from tiny root pieces",
    "Amsterdam set aside €8.2 million to fight it; Ameland and Heusden ask residents to report where it grows",
    "Patches are cheapest to remove when found early",
    "But most people cannot tell knotweed from look-alikes such as bindweed",
], image="examples/test-knotweed.jpg", note="Source: gemeente.nu, 2019, 'Strijdplan Japanse duizendknoop samen met inwoners'. SDG 15, target 15.8: invasive alien species.")

slide("Who it is for (and who not)", [
    "For: residents and park volunteers in the Netherlands who see a suspicious plant and want to know in seconds whether to report it",
    "They have a phone, no botany knowledge, and little time",
    "Not for: municipal staff or contractors deciding whether to remove a plant; that needs an expert on site",
    "Not for: identifying other species",
])

slide("How it works", [
    "1. Take or upload a photo in the app",
    "2. The app sends it to our Roboflow model: the step that decides 'knotweed or not'",
    "3. 70%+ → 'Likely knotweed: don't cut it, report it on waarneming.nl'",
    "   below 70% → 'Maybe: take a closer photo' · otherwise 'Probably not'",
    "Why not a general AI model? We tested one (CLIP): same knotweed found, but 4x more false alarms",
], image="examples/app-screenshot.jpg")

slide("Data and model", [
    "832 Dutch iNaturalist photos, community-confirmed species, Creative Commons licences (credits in sources.csv)",
    "302 knotweed · 530 look-alikes: bindweed, Persicaria, Bistorta, Fallopia, bramble, nettle, ivy",
    "Bug we caught: iNaturalist's search mixed knotweed and bindweed, so ~70 photos had the wrong label",
    "Split by plant, not by photo (587 / 172 / 73), so the test never shows a plant the model already saw",
    "Roboflow classification, ViT model, training photos doubled with flips, rotations and brightness changes",
])

slide("Results on photos it never saw", [
    "73 Dutch test photos: 95% right · 24 of 26 knotweed found · 2 of 47 false alarms",
    "168 new photos from Belgium and Germany, all seasons: 93% right · 77 of 83 found · 5 of 85 false alarms",
    "Spring 93% · summer 93% · autumn/winter 95%: spring shoots are only slightly harder",
    "General AI model (CLIP), same 168 photos: 84% · 77 of 83 found · 21 of 85 false alarms",
    "Our model: 4x fewer false alarms than the general model, which could mean fewer unnecessary checks",
    "Worst mistake: a bindweed called knotweed with 98% confidence",
])

slide("Ethics: what could go wrong, and what we did", [
    "Missed knotweed → patch keeps spreading (the expensive mistake)",
    "False alarm → someone cuts or poisons a harmless plant; cutting real knotweed spreads it",
    "What we did: 70% threshold with a 'maybe' zone; the app never tells you to remove anything, only to report it; tested on a second, independent set",
    "Every answer warns it can be wrong; we report missed knotweed, not just accuracy",
    "Photos go to Roboflow's servers, so don't photograph people or house numbers",
    "Next: test on our own street photos; add winter images",
])

slide("Live demo and what we learned", [
    "Demo: upload a photo → answer in a few seconds",
    "Learned: the labels you scrape can be wrong; check them before training",
    "Learned: a high score can hide the mistakes that matter; count the misses and false alarms",
    "What surprised us: "+"We expected the model to fail on young spring shoots, but when we tested it on 168 new photos from Belgium and Germany it found 27 of 30 spring plants. The real problem was somewhere else: iNaturalist's search had mixed up knotweed and bindweed, so about 70 of our photos had the wrong label before we caught it. Testing taught us more than guessing.",
])

prs.core_properties.last_modified_by = ""
prs.core_properties.author = ""
prs.save("slides.pptx")
print("slides.pptx written")
